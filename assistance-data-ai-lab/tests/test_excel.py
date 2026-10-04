import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from adl.excel_report import build_workbook, kpi_formulas
from adl.kpis import KPIS


@pytest.fixture(scope="module")
def workbook_path(star, quality, tmp_path_factory):
    path = tmp_path_factory.mktemp("excel") / "report.xlsx"
    return build_workbook(star, quality, path)


def test_sheets_and_table(workbook_path, star):
    wb = load_workbook(workbook_path)
    assert wb.sheetnames == ["KPI Summary", "Monthly", "Data Quality", "KPI Definitions", "Cases"]
    cases = wb["Cases"]
    assert "tblCases" in cases.tables
    assert cases.max_row == len(star.fact_case) + 1


def test_kpis_are_live_formulas(workbook_path, star):
    ws = load_workbook(workbook_path)["KPI Summary"]
    expected = kpi_formulas(len(star.fact_case))
    for row, kpi in enumerate(KPIS, start=5):
        assert ws.cell(row=row, column=2).value == expected[kpi.key]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_formulas_match_pipeline_after_recalculation(workbook_path, tmp_path):
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path), str(workbook_path)],
        check=True, capture_output=True, timeout=300,
    )
    ws = load_workbook(tmp_path / workbook_path.name, data_only=True)["KPI Summary"]
    for row in range(5, 5 + len(KPIS)):
        excel_value, pipeline_value = ws.cell(row=row, column=2).value, ws.cell(row=row, column=3).value
        assert excel_value == pytest.approx(pipeline_value, abs=1e-6)
        assert ws.cell(row=row, column=4).value == "OK"
